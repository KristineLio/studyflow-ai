import io
import pytest
from fastapi.testclient import TestClient
from app.main import app


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("STUDYFLOW_DB", str(tmp_path / "test.sqlite3"))
    monkeypatch.delenv("AI_API_KEY", raising=False)
    monkeypatch.delenv("AI_MODEL", raising=False)
    monkeypatch.delenv("ALLOW_DEV_EXPORT", raising=False)
    with TestClient(app) as test_client:
        yield test_client


def test_health_and_demo(client):
    assert client.get("/api/health").json()["ai_configured"] is False
    courses = client.get("/api/courses").json()
    assert len(courses) == 1
    assert len(courses[0]["lectures"]) == 1
    detail = client.get("/api/lectures/1").json()
    assert len(detail["questions"]) == 10
    assert "answer" not in detail["questions"][0]
    assert detail["progress"]["overall"] is None


def test_revision_locked_until_all_stages_finished(client):
    data = client.get("/api/lectures/1").json()
    revision_q = next(q for q in data["questions"] if q["stage"] == "revision")
    assert revision_q["source_excerpt"] == ""
    result = client.post(f"/api/questions/{revision_q['id']}/attempts", json={"answer": "A", "hints_used": 0})
    assert result.status_code == 409
    assert client.post(f"/api/questions/{revision_q['id']}/attempts", json={"answer": "A", "hints_used": 1}).status_code == 422
    assert client.get("/api/lectures/1").json()["progress"]["overall"] is None


def test_full_session_revision_blind_then_revealed(client):
    data = client.get("/api/lectures/1").json()
    prior = [q for q in data["questions"] if q["stage"] != "revision"]
    for q in prior:
        answer = "A" if q["choices"] else "Attempted solution"
        r = client.post(f"/api/questions/{q['id']}/attempts", json={"answer": answer}).json()
        if r["score"] is None:
            assert r["evaluated_by"] == "pending_self_grade"
            assert client.post(f"/api/attempts/{r['attempt_id']}/self-grade", json={"score": 0.5}).status_code == 200
    detail = client.get('/api/lectures/1').json()
    assert all(detail['progress']['stages'][s]['complete'] for s in ('theory','tests','code','problems'))
    revision = [q for q in detail["questions"] if q["stage"] == "revision"]
    first = client.post(f"/api/questions/{revision[0]['id']}/attempts", json={"answer": "A"})
    assert first.status_code == 201
    assert first.json()["score"] is None
    assert first.json()["model_answer"] == ""
    assert first.json()["explanation"] == ""
    mid = client.get('/api/lectures/1').json()
    assert mid['latest_attempts'][str(revision[0]['id'])]['score'] is None
    assert mid['progress']['overall'] is None
    last = client.post(f"/api/questions/{revision[1]['id']}/attempts", json={"answer": "B"})
    assert last.status_code == 201
    assert last.json()["model_answer"]
    complete = client.get('/api/lectures/1').json()
    assert complete['progress']['overall'] is not None
    assert complete['latest_attempts'][str(revision[0]['id'])]['score'] is not None


def test_subjective_answer_requires_explicit_self_grade(client):
    question = next(q for q in client.get("/api/lectures/1").json()["questions"] if q["stage"] == "code")
    r = client.post(f"/api/questions/{question['id']}/attempts", json={"answer": "print('sorted')"})
    assert r.status_code == 201
    assert r.json()["score"] is None
    attempt_id = r.json()["attempt_id"]
    scored = client.post(f"/api/attempts/{attempt_id}/self-grade", json={"score": 0})
    assert scored.status_code == 200
    assert client.post(f"/api/attempts/{attempt_id}/self-grade", json={"score": 1}).status_code == 409


def test_text_upload(client):
    lecture_id = client.get("/api/courses").json()[0]["lectures"][0]["id"]
    result = client.post(f"/api/lectures/{lecture_id}/sources", files={"file": ("lecture.txt", b"Python command line arguments are available from sys.argv. " * 3, "text/plain")})
    assert result.status_code == 201
    source = client.get(f"/api/sources/{result.json()['id']}").json()
    assert "sys.argv" in source["content"]
    assert client.post(f"/api/lectures/{lecture_id}/sources", files={"file": ("bad.jpg", b"bad", "image/jpeg")}).status_code == 415


def test_docx_upload(client):
    from docx import Document
    doc = Document()
    doc.add_heading('Python strings', 0)
    doc.add_paragraph('Command-line arguments are available through sys.argv and represent strings.')
    buf = io.BytesIO()
    doc.save(buf)
    result = client.post('/api/lectures/1/sources', files={'file':('lecture.docx',buf.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
    assert result.status_code == 201
    assert 'sys.argv' in client.get(f"/api/sources/{result.json()['id']}").json()['content']


def test_create_course_and_lecture(client):
    new = client.post("/api/courses", json={"name": "Data Structures", "code": "NETB208"})
    assert new.status_code == 201
    lec = client.post("/api/lectures", json={"course_id": new.json()["id"], "title": "L01 ADTs"})
    assert lec.status_code == 201
    assert client.get(f"/api/lectures/{lec.json()['id']}").json()["questions"] == []


def test_developer_export_disabled_by_default(client, monkeypatch):
    assert client.get('/api/export').status_code == 403
    monkeypatch.setenv('ALLOW_DEV_EXPORT','1')
    assert 'courses' in client.get('/api/export').json()


def test_hint_is_safe_and_revision_is_blocked(client):
    questions = client.get('/api/lectures/1').json()['questions']
    first = next(q for q in questions if q['stage']=='theory')
    hint = client.post(f"/api/questions/{first['id']}/hint",json={'draft_answer':'I am unsure','hint_number':1})
    assert hint.status_code == 200
    assert hint.json()['mode'] == 'general_hint_no_ai'
    revision = next(q for q in questions if q['stage']=='revision')
    assert client.post(f"/api/questions/{revision['id']}/hint",json={}).status_code == 409


def test_pdf_extraction_has_page_references(client):
    import fitz
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((50, 50), "Python lecture: read arguments through sys.argv. " * 3)
    data = pdf.tobytes()
    pdf.close()
    uploaded = client.post('/api/lectures/1/sources', files={'file':('python.pdf',data,'application/pdf')})
    assert uploaded.status_code == 201
    extracted = client.get(f"/api/sources/{uploaded.json()['id']}").json()['content']
    assert '[Page 1]' in extracted
    assert 'sys.argv' in extracted